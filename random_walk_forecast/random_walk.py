"""ランダムウォークによる翌日価格の予測モデル.

対数価格 X_t = log S_t が次のランダムウォークに従うと仮定する.

    X_{t+1} = X_t + mu + sigma * eps_{t+1},   eps_{t+1} ~ i.i.d. N(0, 1)

このとき翌日価格は対数正規分布に従う.

    S_{t+1} = S_t * exp(mu + sigma * eps_{t+1})

    中央値   : S_t * exp(mu)
    期待値   : S_t * exp(mu + sigma^2 / 2)
    (1-a) 区間: S_t * exp(mu -/+ z_{1-a/2} * sigma)

mu = 0 (ドリフトなし) が基本モデル. 比較用にドリフトありも推定する.

使い方:
    python random_walk.py data/prices_2026-10-19.xlsx -o forecast.xlsx
"""

import argparse
from dataclasses import dataclass
from datetime import timedelta
from statistics import NormalDist

import numpy as np
import pandas as pd
from openpyxl.utils import get_column_letter

LEVELS = (0.68, 0.95)


def load_prices(path):
    """Excel を読み込み, 行=日付, 列=アセットの DataFrame を返す.

    1 行目: 曜日, 2 行目: 日付, 3 行目以降: アセット名と価格.
    休場日の '#N/A' は NaN にする.
    """
    raw = pd.read_excel(path, header=None)
    dates = pd.to_datetime(raw.iloc[1, 1:].values)
    body = raw.iloc[2:].set_index(0)
    body.columns = dates
    body = body.apply(pd.to_numeric, errors="coerce")
    df = body.T
    df.index.name = "日付"
    df.columns.name = None
    return df


@dataclass
class RandomWalk:
    """対数価格のランダムウォーク. drift=False なら mu = 0 に固定する."""

    drift: bool = False
    mu: float = 0.0
    sigma: float = np.nan
    last_price: float = np.nan
    n_returns: int = 0

    def fit(self, prices):
        """休場日 (NaN) を除いた価格系列から mu, sigma を推定する."""
        s = pd.Series(prices).dropna().to_numpy(float)
        r = np.diff(np.log(s))
        if len(r) < 2:
            raise ValueError("収益率が 2 個以上必要です")
        self.n_returns = len(r)
        self.last_price = s[-1]
        if self.drift:
            self.mu = r.mean()
            self.sigma = r.std(ddof=1)
        else:
            # mu = 0 が既知なので, 分散は平均を引かずに推定する
            self.mu = 0.0
            self.sigma = np.sqrt(np.mean(r**2))
        return self

    def median(self):
        return self.last_price * np.exp(self.mu)

    def mean(self):
        return self.last_price * np.exp(self.mu + 0.5 * self.sigma**2)

    def interval(self, level):
        z = NormalDist().inv_cdf(0.5 + level / 2)
        return (
            self.last_price * np.exp(self.mu - z * self.sigma),
            self.last_price * np.exp(self.mu + z * self.sigma),
        )

    def simulate(self, n_paths, rng):
        """翌日価格をモンテカルロ法で n_paths 個生成する."""
        eps = rng.standard_normal(n_paths)
        return self.last_price * np.exp(self.mu + self.sigma * eps)


def forecast(df, n_paths, seed):
    """全アセットの翌日予測を作る. 解析解とモンテカルロ法の両方を出す."""
    rng = np.random.default_rng(seed)
    rows = []
    for name in df.columns:
        s = df[name]
        rw = RandomWalk(drift=False).fit(s)
        rwd = RandomWalk(drift=True).fit(s)
        sims = rw.simulate(n_paths, rng)
        row = {
            "アセット": name,
            "最終観測日": s.dropna().index[-1].date(),
            "最新値": rw.last_price,
            "収益率の数": rw.n_returns,
            "σ (日次)": rw.sigma,
            "予想値 (中央値)": rw.median(),
            "期待値": rw.mean(),
        }
        for level in LEVELS:
            lo, hi = rw.interval(level)
            pct = int(level * 100)
            row[f"{pct}%下限"] = lo
            row[f"{pct}%上限"] = hi
        row["MC 中央値"] = np.median(sims)
        row["MC 95%下限"] = np.quantile(sims, 0.025)
        row["MC 95%上限"] = np.quantile(sims, 0.975)
        row["参考: μ (ドリフトあり)"] = rwd.mu
        row["参考: 予想値 (ドリフトあり)"] = rwd.median()
        rows.append(row)
    return pd.DataFrame(rows)


def backtest(df, min_returns=4):
    """1 期先予測のウォークフォワード検証.

    各時点 t で t までのデータだけを使ってモデルを推定し, t+1 の価格を予測する.
    誤差は対数誤差 log(実績 / 予測) で測る.
    """
    records = []
    for name in df.columns:
        s = df[name].dropna()
        for t in range(min_returns, len(s) - 1):
            past = s.iloc[: t + 1]
            actual = s.iloc[t + 1]
            for drift in (False, True):
                rw = RandomWalk(drift=drift).fit(past)
                lo, hi = rw.interval(0.95)
                records.append({
                    "アセット": name,
                    "予測対象日": s.index[t + 1].date(),
                    "モデル": "ドリフトあり" if drift else "ドリフトなし",
                    "予測": rw.median(),
                    "実績": actual,
                    "対数誤差": np.log(actual / rw.median()),
                    "95%区間内": lo <= actual <= hi,
                })
    detail = pd.DataFrame(records)
    summary = (
        detail.groupby(["アセット", "モデル"], sort=False)
        .agg(
            検証回数=("対数誤差", "size"),
            RMSE=("対数誤差", lambda e: np.sqrt(np.mean(e**2))),
            MAE=("対数誤差", lambda e: np.mean(np.abs(e))),
            区間カバー率=("95%区間内", "mean"),
        )
        .reset_index()
    )
    return summary, detail


PCT_COLUMNS = {"σ (日次)", "参考: μ (ドリフトあり)", "RMSE", "MAE", "区間カバー率", "対数誤差"}


def write_sheet(writer, df, sheet):
    df.to_excel(writer, sheet_name=sheet, index=False)
    ws = writer.sheets[sheet]
    ws.freeze_panes = "B2"
    for j, col in enumerate(df.columns, start=1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = max(12, len(str(col)) * 2 + 2)
        if col in PCT_COLUMNS:
            fmt = "0.00%"
        elif pd.api.types.is_float_dtype(df[col]):
            fmt = "#,##0.000"
        else:
            continue
        for cell in ws[letter][1:]:
            cell.number_format = fmt


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("input", help="価格データの Excel ファイル")
    parser.add_argument("-o", "--output", default="forecast.xlsx")
    parser.add_argument("--paths", type=int, default=100_000, help="モンテカルロ法の試行回数")
    parser.add_argument("--seed", type=int, default=0)
    args = parser.parse_args()

    df = load_prices(args.input)
    target = df.index[-1] + timedelta(days=1)
    fc = forecast(df, args.paths, args.seed)
    summary, detail = backtest(df)

    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        write_sheet(writer, fc, "予測")
        write_sheet(writer, summary, "バックテスト")
        write_sheet(writer, detail, "バックテスト明細")

    pd.set_option("display.width", 200)
    print(f"予測対象日: {target.date()}")
    print(fc[["アセット", "最新値", "σ (日次)", "予想値 (中央値)", "95%下限", "95%上限"]]
          .to_string(index=False, float_format=lambda x: f"{x:,.4f}"))
    print()
    print(summary.to_string(index=False, float_format=lambda x: f"{x:.4f}"))
    print(f"\n保存しました: {args.output}")


if __name__ == "__main__":
    main()
